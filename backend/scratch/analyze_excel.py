import sys
import os
import openpyxl

sys.path.append(r'c:\Users\thirs\Downloads\ppl management\PeopleManagement\backend')
from database.connection import execute_query

wb = openpyxl.load_workbook(r'c:\Users\thirs\Downloads\ppl management\college_contacts_mobile.xlsx', data_only=True)
ws = wb['FDP 2026-27 Leads Report']

db_colleges = execute_query('SELECT id, name, college_name, mobile, alt_phone, email, lead_id FROM prospects WHERE prospect_type = \'college_contact\'', fetch='all')

db_by_lid_15 = {}
db_by_name = {}
db_by_college_name = {}
for d in db_colleges:
    lid = (d.get('lead_id') or '').strip()
    if lid:
        db_by_lid_15[lid[:15]] = d
        db_by_lid_15[lid] = d
    if d.get('name'):
        db_by_name[d['name'].strip().lower()] = d
    if d.get('college_name'):
        db_by_college_name[d['college_name'].strip().lower()] = d

print("Checking non-lead-id matches:")
for r in range(2, ws.max_row + 1):
    vals = [ws.cell(row=r, column=c).value for c in range(1, 14)]
    company = str(vals[1]).strip() if vals[1] else ''
    lid = str(vals[12]).strip() if vals[12] else ''
    
    if lid and lid in db_by_lid_15:
        continue
    if company and company.lower() in db_by_name:
        match = db_by_name[company.lower()]
        print(f"Row {r}: Company '{company}' (Excel LeadID: {lid}) -> DB ID {match['id']} (DB LeadID: {match['lead_id']})")
    elif company and company.lower() in db_by_college_name:
        match = db_by_college_name[company.lower()]
        print(f"Row {r}: Company '{company}' (Excel LeadID: {lid}) -> DB ID {match['id']} (DB LeadID: {match['lead_id']})")
