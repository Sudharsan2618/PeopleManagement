import sys
import os
import openpyxl

sys.path.append(r'c:\Users\thirs\Downloads\ppl management\PeopleManagement\backend')
from database.connection import execute_query

wb = openpyxl.load_workbook(r'c:\Users\thirs\Downloads\ppl management\college_contacts_mobile.xlsx', data_only=True)
ws = wb['FDP 2026-27 Leads Report']

db_colleges = execute_query('SELECT * FROM prospects WHERE prospect_type = \'college_contact\' ORDER BY id', fetch='all')

print(f"Total DB colleges: {len(db_colleges)}")
print(f"Total Excel rows: {ws.max_row - 1}")

# Check what each sheet in the Excel represents:
# Notice that Excel has:
# 'FDP 2026-27 Leads Report'
# 'Phone Mapping Audit'
# 'Phone Review Required'
# 'Correction Summary'
# '91 Removal Audit'

print("\nLet's check the columns in FDP 2026-27 Leads Report:")
headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
print(headers)

# Let's check how many Excel rows have a match in DB:
# Priority 1: Lead ID (15-char prefix match)
# Priority 2: College name / Company name exact match (case-insensitive)
db_by_lid_15 = {}
db_by_name = {}
for p in db_colleges:
    lid = (p.get('lead_id') or '').strip()
    if lid:
        db_by_lid_15[lid[:15]] = p
    nm = (p.get('name') or '').strip().lower()
    if nm:
        db_by_name[nm] = p

matched_rows = []
unmatched_rows = []

for r in range(2, ws.max_row + 1):
    vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
    contact_name = vals[0]
    company = vals[1]
    mob = vals[2]
    alt1 = vals[3]
    alt2 = vals[4]
    alt3 = vals[5]
    email = vals[6]
    alt_email = vals[7]
    sec_email = vals[8]
    lead_source = vals[9]
    lead_type = vals[10]
    designation = vals[11]
    lid = str(vals[12]).strip() if vals[12] else None
    
    db_p = None
    match_method = None
    if lid and lid[:15] in db_by_lid_15:
        db_p = db_by_lid_15[lid[:15]]
        match_method = "LEAD_ID"
    elif company and str(company).strip().lower() in db_by_name:
        db_p = db_by_name[str(company).strip().lower()]
        match_method = "COMPANY_NAME"
    elif contact_name and str(contact_name).strip().lower() in db_by_name:
        db_p = db_by_name[str(contact_name).strip().lower()]
        match_method = "CONTACT_NAME"
        
    if db_p:
        matched_rows.append((r, db_p, match_method, vals))
    else:
        unmatched_rows.append((r, vals))

print(f"\nMatched: {len(matched_rows)} / {ws.max_row - 1}")
print(f"Unmatched: {len(unmatched_rows)} / {ws.max_row - 1}")
for ur in unmatched_rows:
    print(f"Unmatched Row {ur[0]}: Name='{ur[1][0]}', Company='{ur[1][1]}', LeadID='{ur[1][12]}'")
