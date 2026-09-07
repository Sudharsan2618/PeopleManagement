import openpyxl
from collections import Counter

wb = openpyxl.load_workbook(r'c:\Users\thirs\Downloads\ppl management\college_contacts_mobile.xlsx', data_only=True)
ws = wb['FDP 2026-27 Leads Report']

mobiles = []
for r in range(2, ws.max_row + 1):
    m = ws.cell(row=r, column=3).value
    if m:
        mobiles.append((r, str(m).strip(), ws.cell(row=r, column=2).value, ws.cell(row=r, column=13).value))

counts = Counter([x[1] for x in mobiles])
dups = {m: c for m, c in counts.items() if c > 1}
print(f"Total duplicate mobile numbers in Excel: {len(dups)}")
for m, c in dups.items():
    print(f"\nMobile {m} appears {c} times:")
    for r, mob, comp, lid in mobiles:
        if mob == m:
            print(f"  Row {r}: {comp} (LeadID: {lid})")
